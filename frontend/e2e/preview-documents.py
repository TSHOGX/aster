"""Bounded synthetic Office fixtures; python-docx 1.2 / openpyxl 3.1 / python-pptx 1.0.

Usage: python preview-documents.py OUTPUT_DIRECTORY
Maximum generated workload: 275,544 cells, 15,599 paragraphs, 200 slides; no user content.
"""
import hashlib
import io
import json
import re
from pathlib import Path
import sys
import zipfile
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
manifest = {}


def canonical(stream):
    # Fix both container and core timestamps (openpyxl updates modified on save).
    output = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(output, 'w') as target:
        for item in sorted(source.namelist()):
            info = zipfile.ZipInfo(item, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            content = source.read(item)
            if item == 'docProps/core.xml':
                content = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)', rb'\g<1>2026-01-01T00:00:00Z\2', content)
            if item.endswith('.xlsx'):
                content = canonical(io.BytesIO(content))
            target.writestr(info, content)
    return output.getvalue()


def save(value, name):
    stream = io.BytesIO()
    value.save(stream)
    data = canonical(stream)
    (root / name).write_bytes(data)
    manifest[name] = {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}


for size in ['small', 'large']:
    doc = Document()
    doc.add_heading('Measured document', 0)
    for i in range(10 if size == 'small' else 1000):
        doc.add_paragraph(f'Paragraph {i}: battery measurements and 中文 text. ' * 3)
    save(doc, f'{size}.docx')
    book = Workbook()
    for i in range(20 if size == 'small' else 2000):
        book.active.append([f'Cell {i}:{j}' for j in range(10)])
    save(book, f'{size}.xlsx')
    slides = Presentation()
    for i in range(2 if size == 'small' else 40):
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = f'Measured slide {i}'
        slide.placeholders[1].text = 'Battery measurements and 中文 text.'
    save(slides, f'{size}.pptx')

for count in [41, 200]:
    slides = Presentation()
    for i in range(count):
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = f'Measured slide {i}'
        slide.placeholders[1].text = 'Battery measurements and 中文 text.'
    save(slides, f'{count}-slides.pptx')

book = Workbook(write_only=True)
sheet = book.create_sheet('Measurements')
for row in range(1, 22963):
    sheet.append([f'Cell {row}:{col}' for col in range(12)])
save(book, 'corpus.xlsx')

# Populate a generated legal OOXML body directly to avoid python-docx's repeated
# table traversal cost; no document bytes come from private evidence.
doc = Document()
stream = io.BytesIO()
doc.save(stream)
namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
paragraphs = ''.join(f'<w:p><w:r><w:t>Paragraph {i}</w:t></w:r></w:p>' for i in range(599))
rows = ''.join('<w:tr>' + ''.join(f'<w:tc><w:p><w:r><w:t>Row {row} cell {col}</w:t></w:r></w:p></w:tc>' for col in range(5)) + '</w:tr>' for row in range(3000))
body = f'<w:document xmlns:w="{namespace}"><w:body>{paragraphs}<w:tbl>{rows}</w:tbl></w:body></w:document>'
output = io.BytesIO()
with zipfile.ZipFile(stream) as source, zipfile.ZipFile(output, 'w') as target:
    for name in source.namelist():
        target.writestr(name, body.encode() if name == 'word/document.xml' else source.read(name))
data = canonical(output)
(root / 'corpus.docx').write_bytes(data)
manifest['corpus.docx'] = {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(), 'paragraphs': 15599, 'table_rows': 3000}

for state in ['charts', 'charts-updated']:
    book = Workbook()
    book.active.append(['Sample', 'Capacity'])
    book.active.append(['A', 100 if state == 'charts' else 105])
    book.active.append(['B', 95])
    chart = BarChart()
    chart.add_data(Reference(book.active, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    book.active.add_chart(chart, 'D2')
    save(book, f'{state}.xlsx')
    slides = Presentation()
    for i in range(2):
        slide = slides.slides.add_slide(slides.slide_layouts[5])
        slide.shapes.title.text = f'{state} slide {i}'
        data = CategoryChartData()
        data.categories = ['A', 'B']
        data.add_series('Capacity', (100 if state == 'charts' else 105, 95))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(7), Inches(4), data)
    save(slides, f'{state}.pptx')
(root / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
