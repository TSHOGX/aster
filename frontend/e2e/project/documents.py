"""Deterministic workload documents for the real browser preview journey."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

root = Path('/workspace')
doc = Document()
doc.add_heading('Reviewed document preview', 0)
doc.add_paragraph('Project document rendering includes 中文 and ordinary text.')
doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Measured value'
doc.save(root / 'preview.docx')
book = Workbook()
sheet = book.active
sheet.append(['Sample', 'Capacity'])
sheet.append(['A', 100])
sheet.append(['B', 95])
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
sheet.add_chart(chart, 'D2')
book.save(root / 'preview.xlsx')
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[5])
slide.shapes.title.text = 'Reviewed presentation preview'
data = CategoryChartData()
data.categories = ['A', 'B']
data.add_series('Capacity', (100, 95))
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(7), Inches(4), data)
slides.save(root / 'preview.pptx')
print('DOCUMENTS-READY')
